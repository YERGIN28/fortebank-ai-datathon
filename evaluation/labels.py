
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
import re

ROOT = Path(__file__).resolve().parents[1]
EXCEL = ROOT / "data" / "open_labels.xlsx"

# Эти пары содержат одинаковые видео.
# Ничего не удаляем — только отмечаем дубликаты.
DUPLICATE_PAIRS = [
    ("X-4", "X-47"),
    ("X-5", "X-48"),
]


def normalize_video(value):
    text = str(value).strip().upper()
    text = text.replace("Х", "X")

    match = re.search(r"X[-_ ]*(\d+)", text)

    if match:
        return f"X-{int(match.group(1))}"

    return text


def parse_timestamp(value):
    if value is None or str(value).strip() == "":
        return None

    text = str(value).strip().lower()

    # Например: "1 минута 8 секунд"
    minutes = re.search(r"(\d+)\s*мин", text)
    seconds = re.search(r"(\d+)\s*сек", text)

    if minutes or seconds:
        return (
            (int(minutes.group(1)) * 60 if minutes else 0)
            + (int(seconds.group(1)) if seconds else 0)
        )

    # Например: "57"
    if re.fullmatch(r"\d+", text):
        return int(text)

    # Например: "01:08"
    match = re.fullmatch(r"(\d+):(\d{2})", text)
    if match:
        return int(match.group(1)) * 60 + int(match.group(2))

    raise ValueError(f"Неизвестный таймкод: {value}")


workbook = load_workbook(
    EXCEL,
    read_only=True,
    data_only=True
)
sheet = workbook.active

records = []

for row in sheet.iter_rows(min_row=2, values_only=True):
    video, timestamp, label = row[:3]

    if not video:
        continue

    records.append({
        "video": normalize_video(video),
        "timestamp_sec": parse_timestamp(timestamp),
        "label": str(label).strip().lower(),
    })

workbook.close()

print("Всего записей:", len(records))
print("Метки:", dict(Counter(r["label"] for r in records)))

print("\nПроверка дубликатов:")

for first, second in DUPLICATE_PAIRS:
    first_records = [
        r for r in records if r["video"] == first
    ]
    second_records = [
        r for r in records if r["video"] == second
    ]

    print(f"\n{first} и {second}:")
    print(f"  {first}: {first_records}")
    print(f"  {second}: {second_records}")

unique_videos = {
    "X-47": "X-4",
    "X-48": "X-5",
}

unique_ids = {
    unique_videos.get(r["video"], r["video"])
    for r in records
}

print("\nУникальных видео:", len(unique_ids))
print("Все исходные записи сохранены.")
